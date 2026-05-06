from pathlib import Path

from openpyxl import load_workbook

import src.synthea_data_profiler_agent as profiler_module
from src.synthea_data_profiler_agent import (
    profile_synthea_files,
    run_synthea_data_profiler_agent,
    write_excel_report,
)


def write_csv(path: Path, content: str) -> None:
    path.write_text(content.strip() + "\n", encoding="utf-8")


def test_profile_synthea_files_infers_column_metrics(tmp_path: Path) -> None:
    write_csv(
        tmp_path / "patients.csv",
        """
        Id,BIRTHDATE,UPDATED,COST,NAME
        b9c610cd-28a6-4636-ccb6-c7a0d2a4cb85,2019-02-17,2019-02-17T05:07:38Z,12.50,Ada
        c1f1fcaa-82fd-d5b7-3544-c8f9708b06a8,2005-07-04,2019-03-24T05:07:38Z,7.00,
        """,
    )

    state = run_synthea_data_profiler_agent(
        synthea_directory=tmp_path,
        report_path=tmp_path / "profile.xlsx",
        execute=False,
    )

    files = state["file_profiles"]
    columns = {profile.column_name: profile for profile in state["column_profiles"]}

    assert files[0].row_count == 2
    assert columns["Id"].inferred_type == "uuid"
    assert columns["BIRTHDATE"].inferred_type == "date"
    assert columns["UPDATED"].inferred_type == "datetime"
    assert columns["COST"].inferred_type == "decimal"
    assert columns["NAME"].blank_count == 1
    assert columns["NAME"].completeness_pct == 50.0


def test_write_excel_report_creates_expected_workbook(tmp_path: Path, monkeypatch) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    report_path = data_dir / "synthea_profile.xlsx"
    monkeypatch.setattr(profiler_module, "DATA_DIR", data_dir)
    write_csv(
        data_dir / "patients.csv",
        """
        Id,NAME
        1,Ada
        2,Grace
        """,
    )

    profiled_state = run_synthea_data_profiler_agent(
        synthea_directory=data_dir,
        report_path=report_path,
        execute=False,
    )
    write_excel_report(
        {
            **profiled_state,
            "report_path": report_path,
            "synthea_directory": data_dir,
        }
    )

    workbook = load_workbook(report_path)
    assert workbook.sheetnames == ["files_overview", "column_profile", "notes"]
    assert workbook["files_overview"]["A2"].value == "patients.csv"
    assert workbook["column_profile"]["C2"].value == "Id"


def test_profile_caps_distinct_counts(tmp_path: Path) -> None:
    write_csv(
        tmp_path / "patients.csv",
        """
        Id
        1
        2
        3
        """,
    )

    state = profile_synthea_files(
        {
            "synthea_files": [tmp_path / "patients.csv"],
            "sample_limit": 2,
            "distinct_limit": 2,
        }
    )

    profile = state["column_profiles"][0]
    assert profile.distinct_count_display == "2+"
    assert profile.distinct_count_is_capped is True
