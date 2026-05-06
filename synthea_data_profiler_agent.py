from __future__ import annotations

import argparse
import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal, TypedDict

from langgraph.graph import END, StateGraph
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
DEFAULT_SYNTHEA_DIR = DATA_DIR / "synthea_tables"
DEFAULT_REPORT_PATH = DATA_DIR / "synthea_data_profile.xlsx"
DEFAULT_SAMPLE_VALUES = 5
DEFAULT_DISTINCT_LIMIT = 1_000

csv.field_size_limit(10_000_000)

UUID_RE = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$"
)
INTEGER_RE = re.compile(r"^[+-]?\d+$")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d+)?Z?$")

BOOLEAN_TRUE_VALUES = {"true", "t", "yes", "y", "1"}
BOOLEAN_FALSE_VALUES = {"false", "f", "no", "n", "0"}


@dataclass(frozen=True)
class FileProfile:
    file_name: str
    row_count: int
    column_count: int
    delimiter: str
    file_size_bytes: int
    blank_cells: int


@dataclass(frozen=True)
class ColumnProfile:
    file_name: str
    ordinal_position: int
    column_name: str
    inferred_type: str
    non_null_count: int
    blank_count: int
    completeness_pct: float
    distinct_count_display: str
    distinct_count_is_capped: bool
    sample_values: tuple[str, ...]
    min_length: int | None
    max_length: int | None
    min_value: str | None
    max_value: str | None


@dataclass
class ColumnAccumulator:
    column_name: str
    ordinal_position: int
    sample_limit: int
    distinct_limit: int
    non_null_count: int = 0
    blank_count: int = 0
    min_length: int | None = None
    max_length: int | None = None
    sample_values: list[str] = field(default_factory=list)
    tracked_distinct_values: set[str] = field(default_factory=set)
    distinct_count_is_capped: bool = False
    possible_boolean: bool = True
    possible_integer: bool = True
    possible_decimal: bool = True
    possible_date: bool = True
    possible_datetime: bool = True
    possible_uuid: bool = True
    numeric_min: Decimal | None = None
    numeric_max: Decimal | None = None
    date_min: date | None = None
    date_max: date | None = None
    datetime_min: datetime | None = None
    datetime_max: datetime | None = None

    def update(self, raw_value: str) -> None:
        value = raw_value.strip()
        if value == "":
            self.blank_count += 1
            return

        self.non_null_count += 1
        value_length = len(value)
        self.min_length = value_length if self.min_length is None else min(self.min_length, value_length)
        self.max_length = value_length if self.max_length is None else max(self.max_length, value_length)

        if value not in self.sample_values and len(self.sample_values) < self.sample_limit:
            self.sample_values.append(value)

        if not self.distinct_count_is_capped:
            if value in self.tracked_distinct_values:
                pass
            elif len(self.tracked_distinct_values) < self.distinct_limit:
                self.tracked_distinct_values.add(value)
            else:
                self.distinct_count_is_capped = True

        normalized_value = value.lower()
        if self.possible_boolean and normalized_value not in BOOLEAN_TRUE_VALUES | BOOLEAN_FALSE_VALUES:
            self.possible_boolean = False

        if self.possible_integer:
            if INTEGER_RE.match(value):
                integer_value = Decimal(int(value))
                self.numeric_min = integer_value if self.numeric_min is None else min(self.numeric_min, integer_value)
                self.numeric_max = integer_value if self.numeric_max is None else max(self.numeric_max, integer_value)
            else:
                self.possible_integer = False

        if self.possible_decimal:
            decimal_value = parse_decimal(value)
            if decimal_value is None:
                self.possible_decimal = False
            else:
                self.numeric_min = decimal_value if self.numeric_min is None else min(self.numeric_min, decimal_value)
                self.numeric_max = decimal_value if self.numeric_max is None else max(self.numeric_max, decimal_value)

        if self.possible_date:
            if DATE_RE.match(value):
                parsed_date = date.fromisoformat(value)
                self.date_min = parsed_date if self.date_min is None else min(self.date_min, parsed_date)
                self.date_max = parsed_date if self.date_max is None else max(self.date_max, parsed_date)
            else:
                self.possible_date = False

        if self.possible_datetime:
            parsed_datetime = parse_datetime(value)
            if parsed_datetime is None:
                self.possible_datetime = False
            else:
                self.datetime_min = (
                    parsed_datetime if self.datetime_min is None else min(self.datetime_min, parsed_datetime)
                )
                self.datetime_max = (
                    parsed_datetime if self.datetime_max is None else max(self.datetime_max, parsed_datetime)
                )

        if self.possible_uuid and not UUID_RE.match(value):
            self.possible_uuid = False

    def to_profile(self, file_name: str, row_count: int) -> ColumnProfile:
        inferred_type = infer_type(self)
        min_value: str | None = None
        max_value: str | None = None

        if inferred_type in {"integer", "decimal"} and self.numeric_min is not None and self.numeric_max is not None:
            min_value = format_decimal(self.numeric_min)
            max_value = format_decimal(self.numeric_max)
        elif inferred_type == "date" and self.date_min is not None and self.date_max is not None:
            min_value = self.date_min.isoformat()
            max_value = self.date_max.isoformat()
        elif inferred_type == "datetime" and self.datetime_min is not None and self.datetime_max is not None:
            min_value = self.datetime_min.isoformat(sep=" ")
            max_value = self.datetime_max.isoformat(sep=" ")

        tracked_count = len(self.tracked_distinct_values)
        distinct_count_display = (
            f"{tracked_count}+"
            if self.distinct_count_is_capped
            else str(tracked_count)
        )
        completeness_pct = (
            round((self.non_null_count / row_count) * 100, 2)
            if row_count
            else 0.0
        )

        return ColumnProfile(
            file_name=file_name,
            ordinal_position=self.ordinal_position,
            column_name=self.column_name,
            inferred_type=inferred_type,
            non_null_count=self.non_null_count,
            blank_count=self.blank_count,
            completeness_pct=completeness_pct,
            distinct_count_display=distinct_count_display,
            distinct_count_is_capped=self.distinct_count_is_capped,
            sample_values=tuple(self.sample_values),
            min_length=self.min_length,
            max_length=self.max_length,
            min_value=min_value,
            max_value=max_value,
        )


class DataProfilerState(TypedDict, total=False):
    synthea_directory: Path
    report_path: Path
    sample_limit: int
    distinct_limit: int
    execute: bool
    synthea_files: list[Path]
    file_profiles: list[FileProfile]
    column_profiles: list[ColumnProfile]
    summary: dict[str, Any]


def parse_decimal(value: str) -> Decimal | None:
    try:
        return Decimal(value)
    except InvalidOperation:
        return None


def parse_datetime(value: str) -> datetime | None:
    if not DATETIME_RE.match(value):
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    except ValueError:
        return None
    if parsed.tzinfo is not None:
        return parsed.astimezone(timezone.utc).replace(tzinfo=None)
    return parsed


def infer_type(accumulator: ColumnAccumulator) -> str:
    if accumulator.non_null_count == 0:
        return "empty"
    if accumulator.possible_boolean:
        return "boolean"
    if accumulator.possible_integer:
        return "integer"
    if accumulator.possible_decimal:
        return "decimal"
    if accumulator.possible_datetime:
        return "datetime"
    if accumulator.possible_date:
        return "date"
    if accumulator.possible_uuid:
        return "uuid"
    return "text"


def format_decimal(value: Decimal) -> str:
    normalized = format(value.normalize(), "f")
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return normalized


def detect_delimiter(file_path: Path) -> str:
    with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t|")
        return dialect.delimiter
    except csv.Error:
        return ","


def validate_report_path(report_path: Path) -> Path:
    resolved = report_path.resolve()
    data_root = DATA_DIR.resolve()
    try:
        resolved.relative_to(data_root)
    except ValueError as error:
        raise ValueError(f"Report path must be inside {data_root}.") from error
    if resolved.suffix.lower() != ".xlsx":
        raise ValueError("Report path must use the .xlsx extension.")
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def discover_synthea_files(state: DataProfilerState) -> DataProfilerState:
    synthea_directory = Path(state.get("synthea_directory", DEFAULT_SYNTHEA_DIR))
    synthea_files = sorted(synthea_directory.glob("*.csv"), key=lambda path: path.name.lower())
    if not synthea_files:
        raise FileNotFoundError(f"No CSV files were found in {synthea_directory}.")
    return {"synthea_files": synthea_files}


def profile_synthea_files(state: DataProfilerState) -> DataProfilerState:
    sample_limit = int(state.get("sample_limit", DEFAULT_SAMPLE_VALUES))
    distinct_limit = int(state.get("distinct_limit", DEFAULT_DISTINCT_LIMIT))
    file_profiles: list[FileProfile] = []
    column_profiles: list[ColumnProfile] = []

    for file_path in state["synthea_files"]:
        delimiter = detect_delimiter(file_path)
        with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle, delimiter=delimiter)
            header = next(reader, None)
            if not header:
                raise ValueError(f"{file_path.name} is empty or missing a header row.")
            if len(set(header)) != len(header):
                raise ValueError(f"{file_path.name} contains duplicate header names.")

            accumulators = [
                ColumnAccumulator(
                    column_name=column_name,
                    ordinal_position=index + 1,
                    sample_limit=sample_limit,
                    distinct_limit=distinct_limit,
                )
                for index, column_name in enumerate(header)
            ]
            row_count = 0

            for row_number, row in enumerate(reader, start=1):
                if len(row) != len(header):
                    raise ValueError(
                        f"{file_path.name} row {row_number} has {len(row)} columns, "
                        f"expected {len(header)}."
                    )
                row_count += 1
                for accumulator, raw_value in zip(accumulators, row, strict=True):
                    accumulator.update(raw_value)

        blank_cells = sum(accumulator.blank_count for accumulator in accumulators)
        file_profiles.append(
            FileProfile(
                file_name=file_path.name,
                row_count=row_count,
                column_count=len(header),
                delimiter=delimiter,
                file_size_bytes=file_path.stat().st_size,
                blank_cells=blank_cells,
            )
        )
        column_profiles.extend(
            accumulator.to_profile(file_path.name, row_count)
            for accumulator in accumulators
        )

    return {"file_profiles": file_profiles, "column_profiles": column_profiles}


def should_write_report(state: DataProfilerState) -> str:
    return "write" if state.get("execute", True) else "summarize"


def autosize_worksheet(worksheet) -> None:
    widths: dict[int, int] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(str(cell.value)) + 2), 60)
    for column_index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def write_excel_report(state: DataProfilerState) -> DataProfilerState:
    report_path = validate_report_path(Path(state.get("report_path", DEFAULT_REPORT_PATH)))
    workbook = Workbook()

    overview_sheet = workbook.active
    overview_sheet.title = "files_overview"
    overview_headers = [
        "file_name",
        "row_count",
        "column_count",
        "delimiter",
        "file_size_bytes",
        "blank_cells",
    ]
    overview_sheet.append(overview_headers)
    for profile in state["file_profiles"]:
        overview_sheet.append(
            [
                profile.file_name,
                profile.row_count,
                profile.column_count,
                profile.delimiter,
                profile.file_size_bytes,
                profile.blank_cells,
            ]
        )

    column_sheet = workbook.create_sheet("column_profile")
    column_headers = [
        "file_name",
        "ordinal_position",
        "column_name",
        "inferred_type",
        "non_null_count",
        "blank_count",
        "completeness_pct",
        "distinct_count",
        "distinct_count_is_capped",
        "sample_values",
        "min_length",
        "max_length",
        "min_value",
        "max_value",
    ]
    column_sheet.append(column_headers)
    for profile in state["column_profiles"]:
        column_sheet.append(
            [
                profile.file_name,
                profile.ordinal_position,
                profile.column_name,
                profile.inferred_type,
                profile.non_null_count,
                profile.blank_count,
                profile.completeness_pct,
                profile.distinct_count_display,
                profile.distinct_count_is_capped,
                " | ".join(profile.sample_values),
                profile.min_length,
                profile.max_length,
                profile.min_value,
                profile.max_value,
            ]
        )

    notes_sheet = workbook.create_sheet("notes")
    notes_sheet.append(["item", "details"])
    notes_sheet.append(["generated_at_utc", datetime.now(timezone.utc).isoformat()])
    notes_sheet.append(["source_directory", str(Path(state.get("synthea_directory", DEFAULT_SYNTHEA_DIR)).resolve())])
    notes_sheet.append(["report_path", str(report_path)])
    notes_sheet.append(["distinct_count_rule", f"Values are tracked exactly up to {state.get('distinct_limit', DEFAULT_DISTINCT_LIMIT)} unique values per column."])
    notes_sheet.append(["sample_values_rule", f"Up to {state.get('sample_limit', DEFAULT_SAMPLE_VALUES)} example values per column are stored."])

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        autosize_worksheet(worksheet)

    workbook.save(report_path)
    return {"report_path": report_path}


def summarize_run(state: DataProfilerState) -> DataProfilerState:
    file_profiles = state.get("file_profiles", [])
    column_profiles = state.get("column_profiles", [])
    return {
        "summary": {
            "files_profiled": len(file_profiles),
            "columns_profiled": len(column_profiles),
            "total_rows": sum(profile.row_count for profile in file_profiles),
            "report_path": str(Path(state.get("report_path", DEFAULT_REPORT_PATH))),
        }
    }


def build_synthea_data_profiler_graph():
    workflow = StateGraph(DataProfilerState)
    workflow.add_node("discover", discover_synthea_files)
    workflow.add_node("profile", profile_synthea_files)
    workflow.add_node("write", write_excel_report)
    workflow.add_node("summarize", summarize_run)

    workflow.set_entry_point("discover")
    workflow.add_edge("discover", "profile")
    workflow.add_conditional_edges(
        "profile",
        should_write_report,
        {"write": "write", "summarize": "summarize"},
    )
    workflow.add_edge("write", "summarize")
    workflow.add_edge("summarize", END)
    return workflow.compile()


synthea_data_profiler_agent = build_synthea_data_profiler_graph()


def run_synthea_data_profiler_agent(
    synthea_directory: Path = DEFAULT_SYNTHEA_DIR,
    report_path: Path = DEFAULT_REPORT_PATH,
    sample_limit: int = DEFAULT_SAMPLE_VALUES,
    distinct_limit: int = DEFAULT_DISTINCT_LIMIT,
    execute: bool = True,
) -> DataProfilerState:
    initial_state: DataProfilerState = {
        "synthea_directory": Path(synthea_directory),
        "report_path": Path(report_path),
        "sample_limit": sample_limit,
        "distinct_limit": distinct_limit,
        "execute": execute,
    }
    return synthea_data_profiler_agent.invoke(initial_state)


def render_summary(state: DataProfilerState) -> str:
    summary = state.get("summary", {})
    return "\n".join(
        [
            f"Files profiled: {summary.get('files_profiled', 0)}",
            f"Columns profiled: {summary.get('columns_profiled', 0)}",
            f"Total rows scanned: {summary.get('total_rows', 0)}",
            f"Report path: {summary.get('report_path', DEFAULT_REPORT_PATH)}",
        ]
    )


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Profile the CSV files in data/synthea_tables and write an Excel report into data."
        )
    )
    parser.add_argument(
        "--synthea-directory",
        type=Path,
        default=DEFAULT_SYNTHEA_DIR,
        help="Directory containing the Synthea CSV files.",
    )
    parser.add_argument(
        "--report-path",
        type=Path,
        default=DEFAULT_REPORT_PATH,
        help="Excel output path inside the data directory.",
    )
    parser.add_argument(
        "--sample-limit",
        type=int,
        default=DEFAULT_SAMPLE_VALUES,
        help="Maximum number of sample values stored per column.",
    )
    parser.add_argument(
        "--distinct-limit",
        type=int,
        default=DEFAULT_DISTINCT_LIMIT,
        help="Maximum number of distinct values tracked exactly per column.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Scan and summarize without writing the Excel workbook.",
    )
    return parser


def main() -> None:
    args = build_argument_parser().parse_args()
    final_state = run_synthea_data_profiler_agent(
        synthea_directory=args.synthea_directory,
        report_path=args.report_path,
        sample_limit=args.sample_limit,
        distinct_limit=args.distinct_limit,
        execute=not args.dry_run,
    )
    print(render_summary(final_state))


if __name__ == "__main__":
    main()
