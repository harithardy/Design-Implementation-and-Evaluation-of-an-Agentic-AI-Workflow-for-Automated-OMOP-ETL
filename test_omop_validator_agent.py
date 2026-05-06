from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

import src.omop_validator_agent as validator_module
from src.omop_validator_agent import (
    build_validation_status,
    run_omop_validator_agent,
    validate_field_mappings,
    write_validation_report,
)


def write_csv(path: Path, content: str) -> None:
    path.write_text(content.strip() + "\n", encoding="utf-8")


def test_build_validation_status_flags_missing_vocabulary() -> None:
    spec = next(
        mapping
        for mapping in validator_module.MAPPING_SPECS
        if mapping.source_file == "conditions.csv"
        and mapping.source_column == "CODE"
        and mapping.target_column == "condition_concept_id"
    )
    status, note = build_validation_status(spec, {"condition_occurrence": {"condition_concept_id"}}, {"SNOMED": 0})

    assert status == "validated_target_missing_vocabulary"
    assert "SNOMED" in note


def test_validate_field_mappings_marks_unmapped_fields() -> None:
    state = validate_field_mappings(
        {
            "source_field_profiles": [
                validator_module.SourceFieldProfile(
                    file_name="payers.csv",
                    column_name="NAME",
                    row_count=2,
                    sample_values=("Medicare",),
                )
            ],
            "omop_schema": {},
            "vocabulary_counts": {},
        }
    )

    row = state["validation_rows"][0]
    assert row.validation_status == "unmapped_source_field"
    assert row.source_file == "payers.csv"


def test_write_validation_report_creates_workbook(tmp_path: Path, monkeypatch) -> None:
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    monkeypatch.setattr(validator_module, "DATA_DIR", data_dir)
    report_path = data_dir / "validator_report.xlsx"
    synthea_dir = data_dir / "synthea_tables"
    synthea_dir.mkdir()
    write_csv(
        synthea_dir / "conditions.csv",
        """
        START,STOP,PATIENT,ENCOUNTER,CODE,DESCRIPTION
        2013-06-24,2013-07-02,p1,e1,10509002,Acute bronchitis (disorder)
        """,
    )

    state = run_omop_validator_agent(
        synthea_directory=synthea_dir,
        report_path=report_path,
        execute=False,
        engine=FakeEngine(),
    )
    write_validation_report({**state, "report_path": report_path, "synthea_directory": synthea_dir})

    workbook = load_workbook(report_path)
    assert workbook.sheetnames == ["file_summary", "field_mappings", "target_coverage", "vocabulary_status", "notes"]
    assert workbook["file_summary"]["A2"].value == "conditions.csv"


class FakeResult:
    def __init__(self, rows):
        self.rows = rows

    def __iter__(self):
        return iter(self.rows)


class FakeConnection:
    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def execute(self, statement, params=None):
        sql = str(statement)
        if "INFORMATION_SCHEMA.COLUMNS" in sql:
            return FakeResult(
                [
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="condition_start_date"),
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="condition_end_date"),
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="person_id"),
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="visit_occurrence_id"),
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="condition_source_value"),
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="condition_source_concept_id"),
                    SimpleNamespace(TABLE_NAME="condition_occurrence", COLUMN_NAME="condition_concept_id"),
                ]
            )
        if "GROUP BY vocabulary_id" in sql:
            return FakeResult([])
        return FakeResult([])


class FakeEngine:
    def connect(self):
        return FakeConnection()
