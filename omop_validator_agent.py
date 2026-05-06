from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal, TypedDict

from langgraph.graph import END, StateGraph
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import text

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT_ROOT / "data"
DEFAULT_SYNTHEA_DIR = DATA_DIR / "synthea_tables"
DEFAULT_REPORT_PATH = DATA_DIR / "omop_mapping_validation_report.xlsx"
DEFAULT_SAMPLE_ROWS = 50
DEFAULT_SAMPLE_VALUES = 3

csv.field_size_limit(10_000_000)


@dataclass(frozen=True)
class SourceFieldProfile:
    file_name: str
    column_name: str
    row_count: int
    sample_values: tuple[str, ...]


@dataclass(frozen=True)
class MappingSpec:
    source_file: str
    source_column: str
    target_table: str | None
    target_column: str | None
    mapping_kind: Literal[
        "direct",
        "derived",
        "foreign_key",
        "concept_lookup",
        "review_required",
        "no_direct",
    ]
    logic_text: str
    applies_when: str | None = None
    source_vocabulary: str | None = None


@dataclass(frozen=True)
class ValidationRow:
    source_file: str
    source_column: str
    sample_values: tuple[str, ...]
    target_table: str | None
    target_column: str | None
    mapping_kind: str
    applies_when: str | None
    source_vocabulary: str | None
    validation_status: str
    logic_text: str
    validation_note: str


@dataclass(frozen=True)
class FileSummary:
    file_name: str
    row_count: int
    source_columns: int
    mapped_fields: int
    review_fields: int
    unmapped_fields: int
    invalid_fields: int


@dataclass(frozen=True)
class TargetCoverage:
    target_table: str
    target_column: str
    source_field_count: int
    source_fields: str


@dataclass(frozen=True)
class VocabularyStatus:
    vocabulary_id: str
    concept_count: int
    used_by_validator: bool
    note: str


class OmopValidatorState(TypedDict, total=False):
    synthea_directory: Path
    report_path: Path
    sample_rows: int
    sample_values: int
    execute: bool
    engine: Any
    synthea_files: list[Path]
    source_field_profiles: list[SourceFieldProfile]
    omop_schema: dict[str, set[str]]
    vocabulary_counts: dict[str, int]
    validation_rows: list[ValidationRow]
    file_summaries: list[FileSummary]
    target_coverage: list[TargetCoverage]
    vocabulary_status: list[VocabularyStatus]
    summary: dict[str, Any]


def mapping_spec(
    source_file: str,
    source_column: str,
    target_table: str | None,
    target_column: str | None,
    mapping_kind: Literal["direct", "derived", "foreign_key", "concept_lookup", "review_required", "no_direct"],
    logic_text: str,
    applies_when: str | None = None,
    source_vocabulary: str | None = None,
) -> MappingSpec:
    return MappingSpec(
        source_file=source_file,
        source_column=source_column,
        target_table=target_table,
        target_column=target_column,
        mapping_kind=mapping_kind,
        logic_text=logic_text,
        applies_when=applies_when,
        source_vocabulary=source_vocabulary,
    )


MAPPING_SPECS: tuple[MappingSpec, ...] = (
    mapping_spec("patients.csv", "Id", "person", "person_source_value", "direct", "Store the source patient identifier in person_source_value."),
    mapping_spec("patients.csv", "Id", "death", "person_id", "foreign_key", "Derive death.person_id from the patient identifier after the person record is created.", applies_when="Only when DEATHDATE is populated."),
    mapping_spec("patients.csv", "BIRTHDATE", "person", "birth_datetime", "derived", "Convert the source birth date into person.birth_datetime."),
    mapping_spec("patients.csv", "BIRTHDATE", "person", "year_of_birth", "derived", "Split BIRTHDATE into the OMOP year_of_birth component."),
    mapping_spec("patients.csv", "BIRTHDATE", "person", "month_of_birth", "derived", "Split BIRTHDATE into the OMOP month_of_birth component."),
    mapping_spec("patients.csv", "BIRTHDATE", "person", "day_of_birth", "derived", "Split BIRTHDATE into the OMOP day_of_birth component."),
    mapping_spec("patients.csv", "DEATHDATE", "death", "death_date", "derived", "Map the source death date into death.death_date."),
    mapping_spec("patients.csv", "DEATHDATE", "death", "death_datetime", "derived", "Map the source death date into death.death_datetime when a timestamp is synthesized or provided."),
    mapping_spec("patients.csv", "GENDER", "person", "gender_source_value", "direct", "Store the original source gender code or label in gender_source_value."),
    mapping_spec("patients.csv", "GENDER", "person", "gender_source_concept_id", "derived", "Translate the source gender label into the appropriate source concept identifier."),
    mapping_spec("patients.csv", "RACE", "person", "race_source_value", "direct", "Store the original race label in race_source_value."),
    mapping_spec("patients.csv", "RACE", "person", "race_source_concept_id", "derived", "Translate the race label into the mapped source concept identifier."),
    mapping_spec("patients.csv", "ETHNICITY", "person", "ethnicity_source_value", "direct", "Store the original ethnicity label in ethnicity_source_value."),
    mapping_spec("patients.csv", "ETHNICITY", "person", "ethnicity_source_concept_id", "derived", "Translate the ethnicity label into the mapped source concept identifier."),
    mapping_spec("patients.csv", "ADDRESS", "location", "address_1", "direct", "Map the street address to location.address_1."),
    mapping_spec("patients.csv", "CITY", "location", "city", "direct", "Map the city value to location.city."),
    mapping_spec("patients.csv", "STATE", "location", "state", "direct", "Map the state value to location.state."),
    mapping_spec("patients.csv", "COUNTY", "location", "county", "direct", "Map the county value to location.county."),
    mapping_spec("patients.csv", "ZIP", "location", "zip", "direct", "Map the postal code to location.zip."),
    mapping_spec("patients.csv", "LAT", "location", "latitude", "direct", "Map the latitude value to location.latitude."),
    mapping_spec("patients.csv", "LON", "location", "longitude", "direct", "Map the longitude value to location.longitude."),
    mapping_spec("organizations.csv", "Id", "care_site", "care_site_source_value", "direct", "Store the organization identifier in care_site_source_value."),
    mapping_spec("organizations.csv", "NAME", "care_site", "care_site_name", "direct", "Map the organization name to care_site_name."),
    mapping_spec("organizations.csv", "ADDRESS", "location", "address_1", "direct", "Map the organization street address to location.address_1."),
    mapping_spec("organizations.csv", "CITY", "location", "city", "direct", "Map the organization city to location.city."),
    mapping_spec("organizations.csv", "STATE", "location", "state", "direct", "Map the organization state to location.state."),
    mapping_spec("organizations.csv", "ZIP", "location", "zip", "direct", "Map the organization postal code to location.zip."),
    mapping_spec("organizations.csv", "LAT", "location", "latitude", "direct", "Map the organization latitude to location.latitude."),
    mapping_spec("organizations.csv", "LON", "location", "longitude", "direct", "Map the organization longitude to location.longitude."),
    mapping_spec("providers.csv", "Id", "provider", "provider_source_value", "direct", "Store the source provider identifier in provider_source_value."),
    mapping_spec("providers.csv", "ORGANIZATION", "provider", "care_site_id", "foreign_key", "Resolve the provider's organization identifier to care_site_id through the care_site ETL."),
    mapping_spec("providers.csv", "ORGANIZATION", "care_site", "care_site_source_value", "direct", "Use the organization identifier as the care site source value."),
    mapping_spec("providers.csv", "NAME", "provider", "provider_name", "direct", "Map the provider display name to provider_name."),
    mapping_spec("providers.csv", "GENDER", "provider", "gender_source_value", "direct", "Store the source provider gender value in gender_source_value."),
    mapping_spec("providers.csv", "GENDER", "provider", "gender_source_concept_id", "derived", "Translate the provider gender label into the OMOP source concept identifier."),
    mapping_spec("providers.csv", "SPECIALITY", "provider", "specialty_source_value", "direct", "Store the source specialty label in specialty_source_value."),
    mapping_spec("providers.csv", "SPECIALITY", "provider", "specialty_source_concept_id", "derived", "Translate the source specialty label into the source concept identifier used by the ETL."),
    mapping_spec("providers.csv", "ADDRESS", "location", "address_1", "direct", "Map the provider practice address to location.address_1."),
    mapping_spec("providers.csv", "CITY", "location", "city", "direct", "Map the provider practice city to location.city."),
    mapping_spec("providers.csv", "STATE", "location", "state", "direct", "Map the provider practice state to location.state."),
    mapping_spec("providers.csv", "ZIP", "location", "zip", "direct", "Map the provider practice postal code to location.zip."),
    mapping_spec("providers.csv", "LAT", "location", "latitude", "direct", "Map the provider practice latitude to location.latitude."),
    mapping_spec("providers.csv", "LON", "location", "longitude", "direct", "Map the provider practice longitude to location.longitude."),
    mapping_spec("encounters.csv", "START", "visit_occurrence", "visit_start_date", "derived", "Map the encounter start timestamp to visit_start_date."),
    mapping_spec("encounters.csv", "START", "visit_occurrence", "visit_start_datetime", "direct", "Map the encounter start timestamp to visit_start_datetime."),
    mapping_spec("encounters.csv", "STOP", "visit_occurrence", "visit_end_date", "derived", "Map the encounter stop timestamp to visit_end_date."),
    mapping_spec("encounters.csv", "STOP", "visit_occurrence", "visit_end_datetime", "direct", "Map the encounter stop timestamp to visit_end_datetime."),
    mapping_spec("encounters.csv", "PATIENT", "visit_occurrence", "person_id", "foreign_key", "Resolve the patient identifier to person_id before loading visit_occurrence."),
    mapping_spec("encounters.csv", "ORGANIZATION", "visit_occurrence", "care_site_id", "foreign_key", "Resolve the organization identifier to care_site_id."),
    mapping_spec("encounters.csv", "PROVIDER", "visit_occurrence", "provider_id", "foreign_key", "Resolve the provider identifier to provider_id."),
    mapping_spec("encounters.csv", "ENCOUNTERCLASS", "visit_occurrence", "visit_source_value", "direct", "Use the encounter class as the original visit source value."),
    mapping_spec("encounters.csv", "ENCOUNTERCLASS", "visit_occurrence", "visit_concept_id", "derived", "Translate encounter class values such as wellness, ambulatory, or inpatient into the corresponding visit concept."),
    mapping_spec("encounters.csv", "REASONCODE", "condition_occurrence", "condition_source_value", "review_required", "Encounter reason codes often become linked condition occurrences, but the ETL must decide whether to persist them as diagnoses or only as visit metadata.", source_vocabulary="SNOMED"),
    mapping_spec("encounters.csv", "REASONCODE", "condition_occurrence", "condition_concept_id", "review_required", "If encounter reasons are materialized as diagnosis facts, translate the SNOMED reason code into a standard condition concept.", source_vocabulary="SNOMED"),
    mapping_spec("conditions.csv", "START", "condition_occurrence", "condition_start_date", "direct", "Map the source condition start date to condition_start_date."),
    mapping_spec("conditions.csv", "STOP", "condition_occurrence", "condition_end_date", "direct", "Map the source condition stop date to condition_end_date."),
    mapping_spec("conditions.csv", "PATIENT", "condition_occurrence", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("conditions.csv", "ENCOUNTER", "condition_occurrence", "visit_occurrence_id", "foreign_key", "Resolve the encounter identifier to visit_occurrence_id."),
    mapping_spec("conditions.csv", "CODE", "condition_occurrence", "condition_source_value", "direct", "Store the source diagnosis code in condition_source_value."),
    mapping_spec("conditions.csv", "CODE", "condition_occurrence", "condition_source_concept_id", "concept_lookup", "Resolve the source SNOMED code to the corresponding source concept identifier.", source_vocabulary="SNOMED"),
    mapping_spec("conditions.csv", "CODE", "condition_occurrence", "condition_concept_id", "concept_lookup", "Resolve the source SNOMED code to the standard condition concept identifier.", source_vocabulary="SNOMED"),
    mapping_spec("procedures.csv", "START", "procedure_occurrence", "procedure_date", "derived", "Map the procedure start timestamp to procedure_date."),
    mapping_spec("procedures.csv", "START", "procedure_occurrence", "procedure_datetime", "direct", "Map the procedure start timestamp to procedure_datetime."),
    mapping_spec("procedures.csv", "STOP", "procedure_occurrence", "procedure_end_date", "derived", "Map the procedure stop timestamp to procedure_end_date."),
    mapping_spec("procedures.csv", "STOP", "procedure_occurrence", "procedure_end_datetime", "direct", "Map the procedure stop timestamp to procedure_end_datetime."),
    mapping_spec("procedures.csv", "PATIENT", "procedure_occurrence", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("procedures.csv", "ENCOUNTER", "procedure_occurrence", "visit_occurrence_id", "foreign_key", "Resolve the encounter identifier to visit_occurrence_id."),
    mapping_spec("procedures.csv", "CODE", "procedure_occurrence", "procedure_source_value", "direct", "Store the source procedure code in procedure_source_value."),
    mapping_spec("procedures.csv", "CODE", "procedure_occurrence", "procedure_source_concept_id", "concept_lookup", "Resolve the source SNOMED procedure code to the source concept identifier.", source_vocabulary="SNOMED"),
    mapping_spec("procedures.csv", "CODE", "procedure_occurrence", "procedure_concept_id", "concept_lookup", "Resolve the source SNOMED procedure code to the standard OMOP procedure concept.", source_vocabulary="SNOMED"),
    mapping_spec("procedures.csv", "BASE_COST", "cost", "total_cost", "direct", "Write the procedure financial amount into a related OMOP cost record with cost_domain_id='Procedure'."),
    mapping_spec("medications.csv", "START", "drug_exposure", "drug_exposure_start_date", "derived", "Map the medication start timestamp to drug_exposure_start_date."),
    mapping_spec("medications.csv", "START", "drug_exposure", "drug_exposure_start_datetime", "direct", "Map the medication start timestamp to drug_exposure_start_datetime."),
    mapping_spec("medications.csv", "STOP", "drug_exposure", "drug_exposure_end_date", "derived", "Map the medication stop timestamp to drug_exposure_end_date."),
    mapping_spec("medications.csv", "STOP", "drug_exposure", "drug_exposure_end_datetime", "direct", "Map the medication stop timestamp to drug_exposure_end_datetime."),
    mapping_spec("medications.csv", "PATIENT", "drug_exposure", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("medications.csv", "ENCOUNTER", "drug_exposure", "visit_occurrence_id", "foreign_key", "Resolve the encounter identifier to visit_occurrence_id."),
    mapping_spec("medications.csv", "CODE", "drug_exposure", "drug_source_value", "direct", "Store the source medication code in drug_source_value."),
    mapping_spec("medications.csv", "CODE", "drug_exposure", "drug_source_concept_id", "concept_lookup", "Resolve the source RxNorm code to the source concept identifier.", source_vocabulary="RxNorm"),
    mapping_spec("medications.csv", "CODE", "drug_exposure", "drug_concept_id", "concept_lookup", "Resolve the source RxNorm code to the standard drug concept identifier.", source_vocabulary="RxNorm"),
    mapping_spec("medications.csv", "DISPENSES", "drug_exposure", "refills", "direct", "Map the dispense count to drug_exposure.refills when the ETL models dispense events as refill counts."),
    mapping_spec("medications.csv", "TOTALCOST", "cost", "total_cost", "direct", "Write the medication total cost into a related OMOP cost record with cost_domain_id='Drug'."),
    mapping_spec("medications.csv", "PAYER_COVERAGE", "cost", "paid_by_payer", "direct", "Write payer coverage into the OMOP cost record as paid_by_payer."),
    mapping_spec("immunizations.csv", "DATE", "drug_exposure", "drug_exposure_start_date", "derived", "Map the immunization date to drug_exposure_start_date."),
    mapping_spec("immunizations.csv", "DATE", "drug_exposure", "drug_exposure_start_datetime", "direct", "Map the immunization date to drug_exposure_start_datetime when a timestamp is retained."),
    mapping_spec("immunizations.csv", "PATIENT", "drug_exposure", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("immunizations.csv", "ENCOUNTER", "drug_exposure", "visit_occurrence_id", "foreign_key", "Resolve the encounter identifier to visit_occurrence_id."),
    mapping_spec("immunizations.csv", "CODE", "drug_exposure", "drug_source_value", "direct", "Store the CVX vaccine code in drug_source_value."),
    mapping_spec("immunizations.csv", "CODE", "drug_exposure", "drug_source_concept_id", "concept_lookup", "Resolve the CVX code to the vaccine source concept identifier.", source_vocabulary="CVX"),
    mapping_spec("immunizations.csv", "CODE", "drug_exposure", "drug_concept_id", "concept_lookup", "Resolve the CVX code to the standard drug concept used for immunizations.", source_vocabulary="CVX"),
    mapping_spec("observations.csv", "DATE", "measurement", "measurement_date", "review_required", "When TYPE='numeric' or the record is routed to Measurement, map DATE to measurement_date.", applies_when="TYPE='numeric'"),
    mapping_spec("observations.csv", "DATE", "measurement", "measurement_datetime", "review_required", "When TYPE='numeric' or the record is routed to Measurement, map DATE to measurement_datetime.", applies_when="TYPE='numeric'"),
    mapping_spec("observations.csv", "DATE", "observation", "observation_date", "review_required", "When the record is not routed to Measurement, map DATE to observation_date.", applies_when="TYPE!='numeric'"),
    mapping_spec("observations.csv", "DATE", "observation", "observation_datetime", "review_required", "When the record is not routed to Measurement, map DATE to observation_datetime.", applies_when="TYPE!='numeric'"),
    mapping_spec("observations.csv", "PATIENT", "measurement", "person_id", "review_required", "Resolve patient to person_id for measurement records.", applies_when="TYPE='numeric'"),
    mapping_spec("observations.csv", "PATIENT", "observation", "person_id", "review_required", "Resolve patient to person_id for non-measurement observation records.", applies_when="TYPE!='numeric'"),
    mapping_spec("observations.csv", "ENCOUNTER", "measurement", "visit_occurrence_id", "review_required", "Resolve encounter to visit_occurrence_id for measurement records.", applies_when="TYPE='numeric'"),
    mapping_spec("observations.csv", "ENCOUNTER", "observation", "visit_occurrence_id", "review_required", "Resolve encounter to visit_occurrence_id for non-measurement observation records.", applies_when="TYPE!='numeric'"),
    mapping_spec("observations.csv", "CODE", "measurement", "measurement_source_value", "review_required", "Store the source observation code in measurement_source_value when the record is routed to Measurement.", applies_when="TYPE='numeric'"),
    mapping_spec("observations.csv", "CODE", "measurement", "measurement_concept_id", "review_required", "Resolve the source LOINC code to the standard measurement concept when TYPE='numeric'.", applies_when="TYPE='numeric'", source_vocabulary="LOINC"),
    mapping_spec("observations.csv", "CODE", "observation", "observation_source_value", "review_required", "Store the source observation code in observation_source_value when the record is routed to Observation.", applies_when="TYPE!='numeric'"),
    mapping_spec("observations.csv", "CODE", "observation", "observation_concept_id", "review_required", "Resolve the source LOINC code to the standard observation concept when TYPE!='numeric'.", applies_when="TYPE!='numeric'", source_vocabulary="LOINC"),
    mapping_spec("observations.csv", "VALUE", "measurement", "value_as_number", "review_required", "Write VALUE into value_as_number when the record is numeric.", applies_when="TYPE='numeric'"),
    mapping_spec("observations.csv", "VALUE", "observation", "value_as_string", "review_required", "Write VALUE into value_as_string when the record is not numeric.", applies_when="TYPE!='numeric'"),
    mapping_spec("observations.csv", "UNITS", "measurement", "unit_source_value", "review_required", "Store the source UCUM unit string in unit_source_value for numeric measurement rows.", applies_when="TYPE='numeric'", source_vocabulary="UCUM"),
    mapping_spec("observations.csv", "UNITS", "measurement", "unit_concept_id", "review_required", "Translate the source UCUM unit string to unit_concept_id for numeric measurement rows.", applies_when="TYPE='numeric'", source_vocabulary="UCUM"),
    mapping_spec("devices.csv", "START", "device_exposure", "device_exposure_start_date", "derived", "Map the device exposure start timestamp to device_exposure_start_date."),
    mapping_spec("devices.csv", "START", "device_exposure", "device_exposure_start_datetime", "direct", "Map the device exposure start timestamp to device_exposure_start_datetime."),
    mapping_spec("devices.csv", "STOP", "device_exposure", "device_exposure_end_date", "derived", "Map the device exposure stop timestamp to device_exposure_end_date."),
    mapping_spec("devices.csv", "STOP", "device_exposure", "device_exposure_end_datetime", "direct", "Map the device exposure stop timestamp to device_exposure_end_datetime."),
    mapping_spec("devices.csv", "PATIENT", "device_exposure", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("devices.csv", "ENCOUNTER", "device_exposure", "visit_occurrence_id", "foreign_key", "Resolve the encounter identifier to visit_occurrence_id."),
    mapping_spec("devices.csv", "CODE", "device_exposure", "device_source_value", "direct", "Store the source device code in device_source_value."),
    mapping_spec("devices.csv", "CODE", "device_exposure", "device_concept_id", "concept_lookup", "Resolve the source SNOMED device code to the OMOP device concept.", source_vocabulary="SNOMED"),
    mapping_spec("devices.csv", "UDI", "device_exposure", "unique_device_id", "direct", "Store the full UDI token in unique_device_id."),
    mapping_spec("supplies.csv", "DATE", "device_exposure", "device_exposure_start_date", "derived", "Map the supply issue date to device_exposure_start_date."),
    mapping_spec("supplies.csv", "PATIENT", "device_exposure", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("supplies.csv", "ENCOUNTER", "device_exposure", "visit_occurrence_id", "foreign_key", "Resolve the encounter identifier to visit_occurrence_id."),
    mapping_spec("supplies.csv", "CODE", "device_exposure", "device_source_value", "direct", "Store the supply code in device_source_value when supplies are modeled in device_exposure."),
    mapping_spec("supplies.csv", "CODE", "device_exposure", "device_concept_id", "concept_lookup", "Resolve the supply SNOMED physical object code to the OMOP device concept.", source_vocabulary="SNOMED"),
    mapping_spec("supplies.csv", "QUANTITY", "device_exposure", "quantity", "direct", "Map the supply quantity to device_exposure.quantity."),
    mapping_spec("payer_transitions.csv", "PATIENT", "payer_plan_period", "person_id", "foreign_key", "Resolve the patient identifier to person_id."),
    mapping_spec("payer_transitions.csv", "START_YEAR", "payer_plan_period", "payer_plan_period_start_date", "direct", "Map the payer transition start timestamp to payer_plan_period_start_date."),
    mapping_spec("payer_transitions.csv", "END_YEAR", "payer_plan_period", "payer_plan_period_end_date", "direct", "Map the payer transition end timestamp to payer_plan_period_end_date."),
    mapping_spec("payer_transitions.csv", "PAYER", "payer_plan_period", "payer_source_value", "direct", "Store the payer identifier as payer_source_value."),
    mapping_spec("payer_transitions.csv", "SECONDARY_PAYER", "payer_plan_period", "sponsor_source_value", "review_required", "Secondary payer data may be modeled as sponsor_source_value, but this requires ETL review because OMOP does not have a dedicated secondary payer field."),
    mapping_spec("allergies.csv", "CODE", "observation", "observation_source_value", "review_required", "OMOP has no dedicated allergy table; many ETLs materialize allergy assertions in Observation and store the source code there.", source_vocabulary="SNOMED"),
    mapping_spec("allergies.csv", "CODE", "observation", "observation_concept_id", "review_required", "If allergies are modeled in Observation, resolve the source allergy code to an appropriate observation concept.", source_vocabulary="SNOMED"),
    mapping_spec("allergies.csv", "PATIENT", "observation", "person_id", "review_required", "Resolve patient to person_id if allergy records are loaded into Observation."),
    mapping_spec("allergies.csv", "ENCOUNTER", "observation", "visit_occurrence_id", "review_required", "Resolve encounter to visit_occurrence_id if allergy records are linked to visits."),
    mapping_spec("careplans.csv", "START", "episode", "episode_start_date", "review_required", "Care plans can be modeled in Episode, but the exact episode concept strategy requires implementation review."),
    mapping_spec("careplans.csv", "STOP", "episode", "episode_end_date", "review_required", "If care plans are modeled in Episode, map STOP to episode_end_date."),
    mapping_spec("careplans.csv", "PATIENT", "episode", "person_id", "review_required", "Resolve the patient identifier to person_id for episode records."),
    mapping_spec("careplans.csv", "CODE", "episode", "episode_source_value", "review_required", "Use the source care plan code as episode_source_value when care plans are represented as episodes.", source_vocabulary="SNOMED"),
    mapping_spec("careplans.csv", "CODE", "episode", "episode_source_concept_id", "review_required", "Resolve the source care plan code to episode_source_concept_id if an Episode-based model is chosen.", source_vocabulary="SNOMED"),
    mapping_spec("imaging_studies.csv", "DATE", "procedure_occurrence", "procedure_date", "review_required", "If imaging studies are represented as procedure_occurrence, map DATE to procedure_date."),
    mapping_spec("imaging_studies.csv", "DATE", "procedure_occurrence", "procedure_datetime", "review_required", "If imaging studies are represented as procedure_occurrence, map DATE to procedure_datetime."),
    mapping_spec("imaging_studies.csv", "PATIENT", "procedure_occurrence", "person_id", "review_required", "Resolve patient to person_id for imaging procedure rows."),
    mapping_spec("imaging_studies.csv", "ENCOUNTER", "procedure_occurrence", "visit_occurrence_id", "review_required", "Resolve encounter to visit_occurrence_id for imaging procedure rows."),
    mapping_spec("imaging_studies.csv", "PROCEDURE_CODE", "procedure_occurrence", "procedure_source_value", "review_required", "Store the imaging procedure source code in procedure_source_value.", source_vocabulary="SNOMED"),
    mapping_spec("imaging_studies.csv", "PROCEDURE_CODE", "procedure_occurrence", "procedure_concept_id", "review_required", "Resolve the imaging procedure SNOMED code to the standard procedure concept.", source_vocabulary="SNOMED"),
)


def get_engine() -> Any:
    try:
        from src.connection import engine
    except ModuleNotFoundError:
        from connection import engine
    return engine


def discover_synthea_files(state: OmopValidatorState) -> OmopValidatorState:
    synthea_directory = Path(state.get("synthea_directory", DEFAULT_SYNTHEA_DIR))
    synthea_files = sorted(synthea_directory.glob("*.csv"), key=lambda path: path.name.lower())
    if not synthea_files:
        raise FileNotFoundError(f"No CSV files were found in {synthea_directory}.")
    return {"synthea_files": synthea_files}


def profile_source_fields(state: OmopValidatorState) -> OmopValidatorState:
    sample_rows = int(state.get("sample_rows", DEFAULT_SAMPLE_ROWS))
    sample_values_limit = int(state.get("sample_values", DEFAULT_SAMPLE_VALUES))
    source_field_profiles: list[SourceFieldProfile] = []

    for file_path in state["synthea_files"]:
        with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames is None:
                raise ValueError(f"{file_path.name} is missing a header row.")
            samples: dict[str, list[str]] = {column: [] for column in reader.fieldnames}
            seen: dict[str, set[str]] = {column: set() for column in reader.fieldnames}
            row_count = 0
            for row in reader:
                row_count += 1
                if row_count > sample_rows:
                    continue
                for column in reader.fieldnames:
                    value = (row.get(column) or "").strip()
                    if not value or value in seen[column] or len(samples[column]) >= sample_values_limit:
                        continue
                    seen[column].add(value)
                    samples[column].append(value)
        for column in reader.fieldnames:
            source_field_profiles.append(SourceFieldProfile(file_path.name, column, row_count, tuple(samples[column])))

    return {"source_field_profiles": source_field_profiles}


def load_omop_reference_data(state: OmopValidatorState) -> OmopValidatorState:
    engine = state.get("engine") or get_engine()
    target_tables = sorted({spec.target_table for spec in MAPPING_SPECS if spec.target_table})
    placeholders = ", ".join(f":table_{index}" for index, _ in enumerate(target_tables))
    params = {f"table_{index}": table for index, table in enumerate(target_tables)}
    omop_schema: dict[str, set[str]] = defaultdict(set)
    vocabulary_counts: dict[str, int] = {}

    with engine.connect() as connection:
        connection.execute(text("USE [AI_OMOP]"))
        rows = connection.execute(
            text(
                f"""
                SELECT TABLE_NAME, COLUMN_NAME
                FROM INFORMATION_SCHEMA.COLUMNS
                WHERE TABLE_SCHEMA = 'dbo'
                  AND TABLE_NAME IN ({placeholders})
                """
            ),
            params,
        )
        for row in rows:
            omop_schema[row.TABLE_NAME].add(row.COLUMN_NAME)

        vocab_rows = connection.execute(
            text(
                """
                SELECT vocabulary_id, COUNT(1) AS concept_count
                FROM [dbo].[concept]
                GROUP BY vocabulary_id
                """
            )
        )
        for row in vocab_rows:
            vocabulary_counts[row.vocabulary_id] = int(row.concept_count)

    return {"omop_schema": dict(omop_schema), "vocabulary_counts": vocabulary_counts}


def build_validation_status(
    spec: MappingSpec,
    omop_schema: dict[str, set[str]],
    vocabulary_counts: dict[str, int],
) -> tuple[str, str]:
    if spec.mapping_kind == "no_direct":
        return "no_direct_omop_field", spec.logic_text

    if not spec.target_table or not spec.target_column:
        return "unmapped_source_field", "No OMOP target column is configured for this mapping rule."

    if spec.target_table not in omop_schema or spec.target_column not in omop_schema[spec.target_table]:
        return "invalid_target_reference", f"Target {spec.target_table}.{spec.target_column} does not exist in AI_OMOP."

    note = f"Validated target {spec.target_table}.{spec.target_column} in AI_OMOP."
    if spec.mapping_kind == "review_required":
        if spec.source_vocabulary:
            count = vocabulary_counts.get(spec.source_vocabulary, 0)
            note += f" Manual review required. Vocabulary {spec.source_vocabulary} has {count} loaded concepts."
        else:
            note += " Manual review required because the source field can branch or requires ETL-specific interpretation."
        return "review_required", note

    if spec.mapping_kind == "concept_lookup" and spec.source_vocabulary:
        count = vocabulary_counts.get(spec.source_vocabulary, 0)
        if count == 0:
            return "validated_target_missing_vocabulary", f"Validated target {spec.target_table}.{spec.target_column}, but source vocabulary {spec.source_vocabulary} is not loaded in dbo.concept."
        note += f" Vocabulary {spec.source_vocabulary} is loaded with {count} concepts."

    return "validated", note


def validate_field_mappings(state: OmopValidatorState) -> OmopValidatorState:
    profiles = state.get("source_field_profiles", [])
    omop_schema = state.get("omop_schema", {})
    vocabulary_counts = state.get("vocabulary_counts", {})
    specs_by_field: dict[tuple[str, str], list[MappingSpec]] = defaultdict(list)
    for spec in MAPPING_SPECS:
        specs_by_field[(spec.source_file.lower(), spec.source_column.lower())].append(spec)

    validation_rows: list[ValidationRow] = []
    profile_index = {(profile.file_name.lower(), profile.column_name.lower()): profile for profile in profiles}

    for profile in profiles:
        field_specs = specs_by_field.get((profile.file_name.lower(), profile.column_name.lower()), [])
        if not field_specs:
            validation_rows.append(
                ValidationRow(
                    source_file=profile.file_name,
                    source_column=profile.column_name,
                    sample_values=profile.sample_values,
                    target_table=None,
                    target_column=None,
                    mapping_kind="unmapped",
                    applies_when=None,
                    source_vocabulary=None,
                    validation_status="unmapped_source_field",
                    logic_text="No OMOP mapping rule is defined for this source field.",
                    validation_note="Field is source-only or still needs ETL design.",
                )
            )
            continue

        for spec in field_specs:
            status, note = build_validation_status(spec, omop_schema, vocabulary_counts)
            validation_rows.append(
                ValidationRow(
                    source_file=profile.file_name,
                    source_column=profile.column_name,
                    sample_values=profile.sample_values,
                    target_table=spec.target_table,
                    target_column=spec.target_column,
                    mapping_kind=spec.mapping_kind,
                    applies_when=spec.applies_when,
                    source_vocabulary=spec.source_vocabulary,
                    validation_status=status,
                    logic_text=spec.logic_text,
                    validation_note=note,
                )
            )

    for spec in MAPPING_SPECS:
        key = (spec.source_file.lower(), spec.source_column.lower())
        if key in profile_index:
            continue
        status, note = build_validation_status(spec, omop_schema, vocabulary_counts)
        validation_rows.append(
            ValidationRow(
                source_file=spec.source_file,
                source_column=spec.source_column,
                sample_values=(),
                target_table=spec.target_table,
                target_column=spec.target_column,
                mapping_kind=spec.mapping_kind,
                applies_when=spec.applies_when,
                source_vocabulary=spec.source_vocabulary,
                validation_status="source_column_missing",
                logic_text=spec.logic_text,
                validation_note=f"Mapping rule references a source column that was not found. {note}",
            )
        )

    file_to_profiles: dict[str, list[SourceFieldProfile]] = defaultdict(list)
    for profile in profiles:
        file_to_profiles[profile.file_name].append(profile)

    file_summaries: list[FileSummary] = []
    for file_name, file_profiles in sorted(file_to_profiles.items()):
        field_statuses: dict[str, set[str]] = defaultdict(set)
        row_count = file_profiles[0].row_count if file_profiles else 0
        for row in validation_rows:
            if row.source_file == file_name:
                field_statuses[row.source_column].add(row.validation_status)
        mapped_fields = sum(1 for statuses in field_statuses.values() if "validated" in statuses)
        review_fields = sum(1 for statuses in field_statuses.values() if "review_required" in statuses)
        unmapped_fields = sum(1 for statuses in field_statuses.values() if statuses <= {"unmapped_source_field"})
        invalid_fields = sum(1 for statuses in field_statuses.values() if "invalid_target_reference" in statuses or "source_column_missing" in statuses)
        file_summaries.append(FileSummary(file_name, row_count, len(file_profiles), mapped_fields, review_fields, unmapped_fields, invalid_fields))

    coverage_groups: dict[tuple[str, str], set[str]] = defaultdict(set)
    for row in validation_rows:
        if not row.target_table or not row.target_column:
            continue
        coverage_groups[(row.target_table, row.target_column)].add(f"{row.source_file}:{row.source_column}")
    target_coverage = [
        TargetCoverage(table, column, len(source_fields), ", ".join(sorted(source_fields)))
        for (table, column), source_fields in sorted(coverage_groups.items())
    ]

    validator_vocabularies = {spec.source_vocabulary for spec in MAPPING_SPECS if spec.source_vocabulary}
    vocabulary_status = [
        VocabularyStatus(
            vocabulary_id=vocabulary_id,
            concept_count=vocabulary_counts.get(vocabulary_id, 0),
            used_by_validator=True,
            note="Vocabulary available for concept lookup." if vocabulary_counts.get(vocabulary_id, 0) > 0 else "Vocabulary missing from dbo.concept.",
        )
        for vocabulary_id in sorted(validator_vocabularies)
    ]

    return {
        "validation_rows": validation_rows,
        "file_summaries": file_summaries,
        "target_coverage": target_coverage,
        "vocabulary_status": vocabulary_status,
    }


def should_write_report(state: OmopValidatorState) -> str:
    return "write" if state.get("execute", True) else "summarize"


def validate_report_path(report_path: Path) -> Path:
    resolved = report_path.resolve()
    data_root = DATA_DIR.resolve()
    try:
        resolved.relative_to(data_root)
    except ValueError as error:
        raise ValueError(f"Report path must be inside {data_root}.") from error
    if resolved.suffix.lower() != ".xlsx":
        raise ValueError("Report path must use the .xlsx extension.")
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def autosize_worksheet(worksheet) -> None:
    widths: dict[int, int] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            widths[cell.column] = min(max(widths.get(cell.column, 0), len(str(cell.value)) + 2), 80)
    for column_index, width in widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = width


def write_validation_report(state: OmopValidatorState) -> OmopValidatorState:
    report_path = validate_report_path(Path(state.get("report_path", DEFAULT_REPORT_PATH)))
    workbook = Workbook()

    summary_sheet = workbook.active
    summary_sheet.title = "file_summary"
    summary_sheet.append(["file_name", "row_count", "source_columns", "mapped_fields", "review_fields", "unmapped_fields", "invalid_fields"])
    for item in state.get("file_summaries", []):
        summary_sheet.append([item.file_name, item.row_count, item.source_columns, item.mapped_fields, item.review_fields, item.unmapped_fields, item.invalid_fields])

    field_sheet = workbook.create_sheet("field_mappings")
    field_sheet.append(["source_file", "source_column", "sample_values", "target_table", "target_column", "mapping_kind", "applies_when", "source_vocabulary", "validation_status", "logic_text", "validation_note"])
    for item in state.get("validation_rows", []):
        field_sheet.append([item.source_file, item.source_column, " | ".join(item.sample_values), item.target_table, item.target_column, item.mapping_kind, item.applies_when, item.source_vocabulary, item.validation_status, item.logic_text, item.validation_note])

    coverage_sheet = workbook.create_sheet("target_coverage")
    coverage_sheet.append(["target_table", "target_column", "source_field_count", "source_fields"])
    for item in state.get("target_coverage", []):
        coverage_sheet.append([item.target_table, item.target_column, item.source_field_count, item.source_fields])

    vocab_sheet = workbook.create_sheet("vocabulary_status")
    vocab_sheet.append(["vocabulary_id", "concept_count", "used_by_validator", "note"])
    for item in state.get("vocabulary_status", []):
        vocab_sheet.append([item.vocabulary_id, item.concept_count, item.used_by_validator, item.note])

    notes_sheet = workbook.create_sheet("notes")
    notes_sheet.append(["item", "details"])
    notes_sheet.append(["generated_at_utc", datetime.now(timezone.utc).isoformat()])
    notes_sheet.append(["report_path", str(report_path)])
    notes_sheet.append(["source_directory", str(Path(state.get("synthea_directory", DEFAULT_SYNTHEA_DIR)).resolve())])
    notes_sheet.append(["method", "The validator checks explicit Synthea-to-OMOP field rules against the live AI_OMOP schema, flags concept lookups whose source vocabularies are missing, and surfaces review-required mappings for ambiguous source domains."])

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        autosize_worksheet(worksheet)

    workbook.save(report_path)
    return {"report_path": report_path}


def summarize_run(state: OmopValidatorState) -> OmopValidatorState:
    validation_rows = state.get("validation_rows", [])
    return {
        "summary": {
            "files_scanned": len(state.get("file_summaries", [])),
            "field_rules_evaluated": len(validation_rows),
            "validated_rows": sum(1 for row in validation_rows if row.validation_status == "validated"),
            "review_rows": sum(1 for row in validation_rows if row.validation_status == "review_required"),
            "unmapped_rows": sum(1 for row in validation_rows if row.validation_status == "unmapped_source_field"),
            "report_path": str(Path(state.get("report_path", DEFAULT_REPORT_PATH))),
        }
    }


def build_omop_validator_graph():
    workflow = StateGraph(OmopValidatorState)
    workflow.add_node("discover", discover_synthea_files)
    workflow.add_node("profile", profile_source_fields)
    workflow.add_node("reference", load_omop_reference_data)
    workflow.add_node("validate", validate_field_mappings)
    workflow.add_node("write", write_validation_report)
    workflow.add_node("summarize", summarize_run)

    workflow.set_entry_point("discover")
    workflow.add_edge("discover", "profile")
    workflow.add_edge("profile", "reference")
    workflow.add_edge("reference", "validate")
    workflow.add_conditional_edges("validate", should_write_report, {"write": "write", "summarize": "summarize"})
    workflow.add_edge("write", "summarize")
    workflow.add_edge("summarize", END)
    return workflow.compile()


omop_validator_agent = build_omop_validator_graph()


def run_omop_validator_agent(
    synthea_directory: Path = DEFAULT_SYNTHEA_DIR,
    report_path: Path = DEFAULT_REPORT_PATH,
    sample_rows: int = DEFAULT_SAMPLE_ROWS,
    sample_values: int = DEFAULT_SAMPLE_VALUES,
    execute: bool = True,
    engine: Any | None = None,
) -> OmopValidatorState:
    initial_state: OmopValidatorState = {
        "synthea_directory": Path(synthea_directory),
        "report_path": Path(report_path),
        "sample_rows": sample_rows,
        "sample_values": sample_values,
        "execute": execute,
    }
    if engine is not None:
        initial_state["engine"] = engine
    return omop_validator_agent.invoke(initial_state)


def render_summary(state: OmopValidatorState) -> str:
    summary = state.get("summary", {})
    return "\n".join([
        f"Files scanned: {summary.get('files_scanned', 0)}",
        f"Field rules evaluated: {summary.get('field_rules_evaluated', 0)}",
        f"Validated rows: {summary.get('validated_rows', 0)}",
        f"Review rows: {summary.get('review_rows', 0)}",
        f"Unmapped rows: {summary.get('unmapped_rows', 0)}",
        f"Report path: {summary.get('report_path', DEFAULT_REPORT_PATH)}",
    ])


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Validate Synthea-to-OMOP field mappings against the live AI_OMOP schema and write an Excel report.")
    parser.add_argument("--synthea-directory", type=Path, default=DEFAULT_SYNTHEA_DIR, help="Directory containing the Synthea CSV files.")
    parser.add_argument("--report-path", type=Path, default=DEFAULT_REPORT_PATH, help="Excel report path inside the data directory.")
    parser.add_argument("--sample-rows", type=int, default=DEFAULT_SAMPLE_ROWS, help="Number of source rows to scan for sample values.")
    parser.add_argument("--sample-values", type=int, default=DEFAULT_SAMPLE_VALUES, help="Maximum unique sample values stored per field.")
    parser.add_argument("--dry-run", action="store_true", help="Validate and summarize without writing the Excel workbook.")
    return parser


def main() -> None:
    args = build_argument_parser().parse_args()
    final_state = run_omop_validator_agent(
        synthea_directory=args.synthea_directory,
        report_path=args.report_path,
        sample_rows=args.sample_rows,
        sample_values=args.sample_values,
        execute=not args.dry_run,
    )
    print(render_summary(final_state))


if __name__ == "__main__":
    main()
